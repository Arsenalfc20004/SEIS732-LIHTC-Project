#!/usr/bin/env python3
"""Aggregate-only P1 profile of the selected HUD archive; no source rows printed.

Requires openpyxl for XLSX input. Run from the repository root:
    python tools/profile_hud_local.py downloads/lihtcpub.zip
Use --accept-new-snapshot only when intentionally establishing a new baseline.
"""

import argparse
import csv
import hashlib
import io
import re
import sys
from collections import Counter
from datetime import datetime, timezone
from decimal import Decimal, InvalidOperation
from pathlib import Path
from zipfile import BadZipFile, ZipFile

PINNED_SHA256 = "e07acee706174b276f89596d614ac5699efa9848659e5834fdfb5198fa0a7288"
EXPECTED_COLUMNS = """
hud_id project proj_add proj_cty proj_st proj_zip state_id latitude longitude
place1990 place2000 place2010 place2020 fips1990 fips2000 fips2010 fips2020
st2020 cnty2020 trct2020 scattered_site_cd resyndication_cd allocamt n_units
li_units n_0br n_1br n_2br n_3br n_4br aff_period aff_yrs inc_ceil low_ceil
ceilunit yr_pis yr_alloc non_prof basis bond mff_ra fmha_514 fmha_515 fmha_538
home home_amt tcap tcap_amt cdbg cdbg_amt htf htf_amt fha hopevi hpvi_amt tcep
tcep_amt rad qozf qozf_amt rentassist trgt_pop trgt_fam trgt_eld trgt_dis
trgt_hml trgt_oth trgt_spc type credit n_unitsr li_unitr metro dda qct nonprog
nlm_reason nlm_spc datanote record_stat
""".split()

GRAIN_REFERENCE = {
    "NATIONWIDE": (55345, 0, 55345, 0, 0),
    "MINNESOTA": (1144, 0, 1144, 0, 0),
    "MN_VALID_SERVICE_YEAR": (1143, 0, 1143, 0, 0),
}
GRAIN_FIELDS = (
    "source_rows", "rows_with_missing_id", "distinct_nonblank_ids",
    "repeated_id_groups", "excess_rows_in_repeated_groups",
)
QUALITY_REFERENCE = {
    "valid_year_records": 1143,
    "year_9999_records": 1,
    "year_8888_records": 0,
    "missing_or_invalid_year_records": 0,
    "other_outside_window_records": 0,
    "first_valid_service_year": 1987,
    "last_valid_service_year": 2024,
    "blank_city_records": 0,
    "distinct_normalized_city_labels": 266,
    "missing_or_invalid_total_unit_records": 0,
    "missing_or_invalid_low_income_unit_records": 0,
    "inconsistent_unit_records": 2,
    "flagged_resyndication_records": 34,
    "unreported_resyndication_records": 681,
}


def text_value(value):
    """Match an ordinary spreadsheet CSV export without coercing text IDs."""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def integer_value(value):
    text = text_value(value)
    return int(text) if re.fullmatch(r"[+-]?\d+", text) else None


def number_value(value):
    text = text_value(value)
    if not text:
        return None
    try:
        number = Decimal(text)
        return number if number.is_finite() else None
    except InvalidOperation:
        return None


def file_hash(path):
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def grain_counts(counter):
    return (
        sum(counter.values()),
        counter.get(None, 0),
        sum(key is not None for key in counter),
        sum(key is not None and count > 1 for key, count in counter.items()),
        sum(count - 1 for key, count in counter.items()
            if key is not None and count > 1),
    )


def profile_rows(rows):
    header = next(rows, None)
    if header is None:
        raise ValueError("Selected member has no header.")
    columns = [text_value(value).lstrip("\ufeff").lower() for value in header]
    if columns != EXPECTED_COLUMNS:
        raise ValueError(
            "Header does not match the 80 pinned source columns in order. "
            "Inspect the actual schema before updating this check."
        )
    positions = {name: index for index, name in enumerate(columns)}
    counters = {name: Counter() for name in GRAIN_REFERENCE}
    quality = Counter()
    cities = set()
    valid_years = set()
    for row in rows:
        if len(row) == 0:
            quality["empty_physical_lines_skipped"] += 1
            continue
        if len(row) != len(columns):
            raise ValueError("A source record does not have exactly 80 fields.")
        if all(not text_value(value) for value in row):
            quality["completely_empty_source_records"] += 1
        get = lambda field: row[positions[field]]
        project_id = text_value(get("hud_id")) or None
        counters["NATIONWIDE"][project_id] += 1
        if text_value(get("proj_st")).upper() != "MN":
            continue
        counters["MINNESOTA"][project_id] += 1
        service_year = integer_value(get("yr_pis"))
        valid_year = service_year is not None and 1987 <= service_year <= 2024
        if valid_year:
            counters["MN_VALID_SERVICE_YEAR"][project_id] += 1
            quality["valid_year_records"] += 1
            valid_years.add(service_year)
        elif service_year == 9999:
            quality["year_9999_records"] += 1
        elif service_year == 8888:
            quality["year_8888_records"] += 1
        elif service_year is None:
            quality["missing_or_invalid_year_records"] += 1
        else:
            quality["other_outside_window_records"] += 1
        city = text_value(get("proj_cty")).upper()
        if city:
            cities.add(city)
        else:
            quality["blank_city_records"] += 1
        total_units = number_value(get("n_units"))
        low_units = number_value(get("li_units"))
        quality["missing_or_invalid_total_unit_records"] += total_units is None
        quality["missing_or_invalid_low_income_unit_records"] += low_units is None
        if total_units is not None and low_units is not None:
            quality["inconsistent_unit_records"] += low_units > total_units
        units_eligible = (
            total_units is not None and low_units is not None
            and total_units > 0 and 0 <= low_units <= total_units
        )
        quality["all_year_unit_ineligible_records"] += not units_eligible
        quality["valid_year_unit_ineligible_records"] += valid_year and not units_eligible
        flag = text_value(get("resyndication_cd"))
        quality["flagged_resyndication_records"] += flag == "1"
        quality["unreported_resyndication_records"] += not flag
    quality["distinct_normalized_city_labels"] = len(cities)
    quality["first_valid_service_year"] = min(valid_years) if valid_years else None
    quality["last_valid_service_year"] = max(valid_years) if valid_years else None
    return {scope: grain_counts(counter) for scope, counter in counters.items()}, quality


def read_member(archive, member_name, sheet_name):
    suffix = Path(member_name).suffix.lower()
    if suffix == ".xlsx":
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise ValueError("Install the reader with: python -m pip install openpyxl") from exc
        workbook = load_workbook(io.BytesIO(archive.read(member_name)),
                                 read_only=True, data_only=True)
        try:
            if sheet_name not in workbook.sheetnames:
                raise ValueError("Selected worksheet is not in the workbook.")
            return profile_rows(iter(workbook[sheet_name].iter_rows(values_only=True)))
        finally:
            workbook.close()
    if suffix == ".csv":
        with archive.open(member_name) as binary:
            with io.TextIOWrapper(binary, encoding="utf-8-sig", newline="") as stream:
                return profile_rows(iter(csv.reader(stream)))
    raise ValueError("Select the actual .xlsx or .csv project member in the archive.")


def report(input_path, digest, member, sheet, grain, quality):
    pinned = digest == PINNED_SHA256
    lines = [
        "# HUD P1 local aggregate check", "",
        f"- Run timestamp (UTC): {datetime.now(timezone.utc).isoformat()}",
        f"- Input archive: `{input_path.name}`",
        f"- Archive SHA-256: `{digest}`",
        f"- Selected member: `{member}`",
        f"- Worksheet: `{sheet}`" if member.lower().endswith(".xlsx")
        else "- Worksheet: not applicable to CSV",
        "- Header: 80 expected source columns in order (case-insensitive check)",
        "- Source identity: " + ("matches the recorded P0 archive checksum"
                                   if pinned else "different archive; new baseline required"),
        "", "## Grain check", "",
        "| Scope | Check | P0 reference | Observed now | Comparison |",
        "| --- | --- | ---: | ---: | --- |",
    ]
    for scope, values in grain.items():
        for field, expected, observed in zip(GRAIN_FIELDS, GRAIN_REFERENCE[scope], values):
            comparison = ("MATCH" if observed == expected else "REVIEW") if pinned else "NEW INPUT"
            lines.append(f"| {scope} | {field} | {expected} | {observed} | {comparison} |")
    lines += ["", "## Minnesota quality checks", "",
              "| Check | P0 reference | Observed now | Comparison |",
              "| --- | ---: | ---: | --- |"]
    for field, expected in QUALITY_REFERENCE.items():
        observed = quality.get(field, 0)
        comparison = ("MATCH" if observed == expected else "REVIEW") if pinned else "NEW INPUT"
        lines.append(f"| {field} | {expected} | {observed} | {comparison} |")
    for field in ("all_year_unit_ineligible_records", "valid_year_unit_ineligible_records"):
        lines.append(f"| {field} | Not previously recorded | {quality.get(field, 0)} | Measured now |")
    lines += ["", f"Empty physical lines skipped: {quality.get('empty_physical_lines_skipped', 0)}.",
              f"Completely empty 80-field source records counted: {quality.get('completely_empty_source_records', 0)}.",
              "", "These checks concern project records in one release. They do not establish unique physical housing or current affordability."]
    return "\n".join(lines) + "\n"


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("archive", type=Path, help="Actual local HUD ZIP; do not commit it to Git")
    parser.add_argument("--member", default="LIHTCPUB.xlsx", help="Selected project member's basename or full ZIP path")
    parser.add_argument("--sheet", default="Data", help="Worksheet for XLSX input")
    parser.add_argument("--accept-new-snapshot", action="store_true",
                        help="Profile a different archive and label it as a new baseline")
    args = parser.parse_args(argv)
    try:
        digest = file_hash(args.archive)
        if digest != PINNED_SHA256 and not args.accept_new_snapshot:
            raise ValueError(
                "Archive checksum differs from the recorded P0 input. "
                "Confirm the source; use --accept-new-snapshot only to record a deliberate new baseline."
            )
        with ZipFile(args.archive) as archive:
            names = [name for name in archive.namelist() if not name.endswith("/")]
            exact = [name for name in names if name.casefold() == args.member.casefold()]
            matches = exact or [name for name in names
                                if Path(name).name.casefold() == args.member.casefold()]
            if len(matches) != 1:
                raise ValueError("Selected member was not found uniquely; inspect the archive and specify --member.")
            member = matches[0]
            grain, quality = read_member(archive, member, args.sheet)
        print(report(args.archive, digest, member, args.sheet, grain, quality), end="")
    except (OSError, ValueError, BadZipFile, csv.Error) as exc:
        print(f"Profile not completed: {exc}", file=sys.stderr)
        return 2
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

